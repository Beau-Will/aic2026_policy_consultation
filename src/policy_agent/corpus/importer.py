"""语料导入：支持 md / txt / html / pdf / docx → 归一化文本 + 元数据猜测。"""

from __future__ import annotations

import re
from pathlib import Path

from .models import CorpusDocument, sha256_text
from .split import normalize_text, split_articles

SUPPORTED_SUFFIXES = (".md", ".markdown", ".txt", ".html", ".htm", ".pdf", ".docx")

# 常见但当前不支持：扫描目录时给出提示，避免"文件被静默忽略"
UNSUPPORTED_DOC_SUFFIXES = (".doc", ".wps", ".rtf", ".odt", ".mht", ".mhtml")

_RE_DOC_NUMBER = re.compile(
    r"[\u4e00-\u9fa5A-Za-z]{0,12}[〔\[（(](?:19|20)\d{2}[〕\]）)]\s*第?\s*\d+\s*号"
)
# 国务院令 / 部委令 形式的发文字号，例如“国务院令第775号”
_RE_ORDER_NUMBER = re.compile(
    r"((?:中华人民共和国)?国务院令\s*第\s*\d+\s*号|国令第\s*\d+\s*号|(?:[\u4e00-\u9fa5]{2,10}部令)\s*第\s*\d+\s*号)"
)
_RE_DATE = re.compile(r"((?:19|20)\d{2})\s*[年./-]\s*(\d{1,2})\s*[月./-]\s*(\d{1,2})\s*日?")
_RE_ISSUER = re.compile(
    r"^[\u4e00-\u9fa5]{2,30}(?:部|委|局|厅|署|办|办公室|政府|人民法院|人民检察院)$"
)

# 施行/生效：只有原文明确写了“自XXXX年X月X日起施行/生效/实施”才采信
_RE_EFFECTIVE_CLAUSE = re.compile(
    r"自\s*((?:19|20)\d{2})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日起\s*(?:施行|生效|实施)"
)
# 发布/通过：优先取“通过/发布/公布/签发”附近的日期
_RE_PUBLISH_KEYWORD = re.compile(r"(通过|发布|公布|签发|批准)")


def iter_supported_files(root: Path) -> list[Path]:
    """递归列出可导入的原文文件。

    以 "_" 或 "." 开头的文件会被跳过，方便在语料目录里放说明/索引文件。
    """
    return sorted(
        p
        for p in root.rglob("*")
        if p.is_file()
        and p.suffix.lower() in SUPPORTED_SUFFIXES
        and not p.name.startswith(("_", "."))
    )


def _fmt_date(m: re.Match[str]) -> str:
    return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"


_RE_MARKER_LINE = re.compile(r"^(第[一二三四五六七八九十百零〇两0-9]+[章节条]|[一二三四五六七八九十]+[、.．])")


def _looks_like_marker(line: str) -> bool:
    return bool(_RE_MARKER_LINE.match(line))


def _read_md_txt(path: Path) -> tuple[str, str | None]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    title = None
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            title = stripped.lstrip("#").strip()
            break
        # 跳过章节/条号行，避免把“第一章 总则”误当成标题
        if _looks_like_marker(stripped):
            continue
        title = stripped
        break
    return text, title


def _read_html(path: Path) -> tuple[str, str | None]:
    from bs4 import BeautifulSoup

    raw = path.read_text(encoding="utf-8", errors="ignore")
    soup = BeautifulSoup(raw, "lxml")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    title = soup.title.get_text(strip=True) if soup.title else None
    text = soup.get_text("\n")
    return text, title


def _read_pdf(path: Path) -> tuple[str, str | None]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    meta_title = None
    try:
        if reader.metadata and reader.metadata.title:
            meta_title = str(reader.metadata.title).strip() or None
    except Exception:
        meta_title = None
    return "\n\n".join(pages), meta_title


def _read_docx(path: Path) -> tuple[str, str | None]:
    import docx

    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs]
    text = "\n".join(paragraphs)
    title = next((p.strip() for p in paragraphs if p.strip()), None)
    return text, title


def load_text(path: Path) -> tuple[str, str | None]:
    """按扩展名读取文件，返回 (原文文本, 猜测的标题)。"""
    suffix = path.suffix.lower()
    if suffix in (".md", ".markdown", ".txt"):
        return _read_md_txt(path)
    if suffix in (".html", ".htm"):
        return _read_html(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    raise ValueError(f"不支持的文件格式：{suffix}（支持 {'/'.join(SUPPORTED_SUFFIXES)}）")


def guess_metadata(text: str) -> dict[str, str | None]:
    """从原文正文中尽力猜测发文字号、发文机关与日期。"""
    result: dict[str, str | None] = {
        "doc_number": None,
        "issuer": None,
        "publish_date": None,
        "effective_date": None,
    }
    m = _RE_DOC_NUMBER.search(text)
    if m is None:
        m = _RE_ORDER_NUMBER.search(text)
    if m:
        result["doc_number"] = re.sub(r"\s+", "", m.group(0))

    for line in text.splitlines()[:40]:
        stripped = line.strip()
        if _RE_ISSUER.match(stripped):
            result["issuer"] = stripped
            break

    # 发布日期：优先取“通过/发布/公布”关键词附近的日期；否则取文首（前200字）的第一个日期
    head = text[:600]
    keyword = _RE_PUBLISH_KEYWORD.search(head)
    if keyword:
        window = head[max(0, keyword.start() - 60) : keyword.end() + 60]
        found = _RE_DATE.search(window)
        if found:
            result["publish_date"] = _fmt_date(found)
    if result["publish_date"] is None:
        head_date = _RE_DATE.search(text[:200])
        if head_date:
            result["publish_date"] = _fmt_date(head_date)

    # 施行日期：只认明确的“自…起施行”条款；宁可为空，也不猜
    effective = _RE_EFFECTIVE_CLAUSE.search(text)
    if effective:
        result["effective_date"] = (
            f"{int(effective.group(1)):04d}-{int(effective.group(2)):02d}-{int(effective.group(3)):02d}"
        )
    return result


def _slugify(name: str) -> str:
    slug = re.sub(r"[^\w\u4e00-\u9fa5-]+", "-", name).strip("-")
    return slug[:60] or "doc"


def build_document(
    path: Path,
    *,
    doc_id: str | None = None,
    title: str | None = None,
    issuer: str | None = None,
    doc_number: str | None = None,
    publish_date: str | None = None,
    effective_date: str | None = None,
    region: str | None = None,
    source_url: str | None = None,
) -> CorpusDocument:
    """读取文件并构建条文级语料文档。显式传入的元数据优先于自动猜测。"""
    text, guessed_title = load_text(path)
    text = normalize_text(text)
    if not text:
        raise ValueError(f"文件内容为空或无法解析：{path}")

    guessed = guess_metadata(text)
    content_hash = sha256_text(text)
    # doc_id 由“文件名 + 文件路径”派生：
    #   - 同一路径重新导入 -> 同一个 doc_id，覆盖旧版本（便于政策更新）；
    #   - 不同目录下同名文件（如两份“通知.pdf”）-> 不同 doc_id，不会互相覆盖。
    if doc_id is None:
        resolved = str(path.resolve()).replace("\\", "/").lower()
        doc_id = f"{_slugify(path.stem)}-{sha256_text(resolved)[:8]}"

    articles = split_articles(text)
    for index, article in enumerate(articles):
        article.doc_id = doc_id
        article.article_id = f"{doc_id}#{index:04d}"

    return CorpusDocument(
        doc_id=doc_id,
        title=(title or guessed_title or path.stem).strip(),
        raw_text=text,
        source_hash=content_hash,
        issuer=issuer or guessed["issuer"],
        doc_number=doc_number or guessed["doc_number"],
        publish_date=publish_date or guessed["publish_date"],
        effective_date=effective_date or guessed["effective_date"],
        region=region,
        source_url=source_url,
        articles=articles,
    )
