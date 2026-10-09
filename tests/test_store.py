"""语料库导入与检索测试。"""

from __future__ import annotations

import pathlib

import pytest

from policy_agent.corpus.importer import build_document, guess_metadata, load_text
from policy_agent.corpus.store import CorpusStore, normalize_for_match


SAMPLE_NOTICE = """XX市人力资源和社会保障局关于做好2024年度创业补贴申报工作的通知
X人社发〔2024〕7号
2024年3月5日

一、申报对象
在本市登记注册并正常经营满6个月的小微企业、个体工商户。

二、申报时间
申报期为2024年4月1日至4月30日，逾期不再受理。

三、申报材料
（一）营业执照复印件；
（二）近6个月社保缴纳证明。
"""


def test_build_document_md_metadata(sample_md):
    doc = build_document(sample_md)
    assert doc.title.startswith("示例市人民政府办公室")
    assert doc.doc_number == "示政办发〔2024〕12号"
    assert doc.publish_date == "2024-03-01"
    assert len(doc.articles) == 4
    assert doc.articles[0].article_id.startswith(doc.doc_id)


def test_build_document_enum_fallback(tmp_path):
    path = tmp_path / "notice.md"
    path.write_text(SAMPLE_NOTICE, encoding="utf-8")
    doc = build_document(path)
    assert doc.doc_number == "X人社发〔2024〕7号"
    numbers = [a.article_no for a in doc.articles]
    assert numbers == ["一", "二", "三"]


def test_build_document_docx(tmp_path):
    import docx

    path = tmp_path / "policy.docx"
    document = docx.Document()
    document.add_paragraph("某部关于印发管理办法的通知")
    document.add_paragraph("某发〔2024〕3号")
    document.add_paragraph("第一条 为规范管理，制定本办法。")
    document.add_paragraph("第二条 本办法自2024年1月1日起施行。")
    document.save(str(path))

    doc = build_document(path)
    assert "管理办法" in doc.title
    assert len(doc.articles) == 2
    assert doc.articles[0].article_no == "第一条"


def test_build_document_html(tmp_path):
    path = tmp_path / "policy.html"
    path.write_text(
        "<html><head><title>关于开展试点的通知</title></head><body>"
        "<p>一、试点范围</p><p>本市全域。</p>"
        "<p>二、试点时间</p><p>2024年6月1日起。</p>"
        "</body></html>",
        encoding="utf-8",
    )
    doc = build_document(path)
    assert doc.title == "关于开展试点的通知"
    assert len(doc.articles) == 2


def test_load_text_rejects_unsupported(tmp_path):
    path = tmp_path / "a.bin"
    path.write_bytes(b"\x00\x01")
    with pytest.raises(ValueError):
        load_text(path)


def test_guess_metadata_none_when_absent():
    meta = guess_metadata("一段没有任何发文字号与日期的普通文字。")
    assert meta["doc_number"] is None
    assert meta["publish_date"] is None


def test_search_ranks_target_article(populated_store):
    hits = populated_store.search("申报时间 4月30日", limit=3)
    assert hits, "应至少命中一条"
    assert hits[0]["article_no"] == "第四条"


def test_search_empty_query_returns_nothing(populated_store):
    assert populated_store.search("   ") == []


def test_find_quote_hits_and_misses(populated_store):
    assert populated_store.find_quote("申报时间为每年4月1日至4月30日")
    assert populated_store.find_quote("这句话在语料库里根本不存在") == []


def test_reimport_same_doc_replaces_not_duplicates(tmp_path):
    store = CorpusStore(tmp_path / "db.sqlite")
    path = tmp_path / "p.md"
    path.write_text("第一条 内容甲。", encoding="utf-8")
    store.add_document(build_document(path))
    path.write_text("第一条 内容乙。", encoding="utf-8")
    store.add_document(build_document(path))
    stats = store.stats()
    assert stats["documents"] == 1
    assert stats["articles"] == 1
    store.close()


def test_documents_scope(populated_store):
    hits = populated_store.search("小微", limit=1)
    scope = populated_store.documents_scope([hits[0]["doc_id"]])
    assert scope and "示例市" in scope[0]


def test_normalize_for_match_ignores_punctuation_and_space():
    assert normalize_for_match("申报时间： 4月1日。") == normalize_for_match("申报时间4月1日")
